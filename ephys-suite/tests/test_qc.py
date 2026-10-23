import os
import shutil

import numpy as np
import pytest

from ephys_suite import qc
from ephys_suite.qc import QCParams, SweepRa, FileQC


def vc_sweep(fs=20000.0, dur_s=2.0, onset_ms=156.3, pulse_ms=100.0, dv=-10.0, ra=15.0, rm=300.0, cm=30.0, hold=-70.0,
             noise=3.0, epsc=0, seed=0, command=True, i_hold=-9.0):
    """Voltage-clamp sweep of a cell (Ra, Rm in MOhm, Cm in pF) given a voltage step; returns (t_ms, I_pA, V_mV|None)."""
    rng = np.random.default_rng(seed)
    n = int(dur_s * fs)
    t = np.arange(n) / fs * 1e3
    tau_ms = (ra * rm / (ra + rm)) * cm / 1000.0
    i_ss = dv / (ra + rm) * 1000.0
    i_0 = dv / ra * 1000.0
    I = np.full(n, i_hold)
    on = int(round(onset_ms / 1e3 * fs))
    off = on + int(round(pulse_ms / 1e3 * fs))
    x = (np.arange(n) - on) / fs * 1e3
    I[on:off] += i_ss + (i_0 - i_ss) * np.exp(-x[on:off] / tau_ms)
    # offset: the membrane (charged to dv * Rm / (Ra + Rm)) discharges through Ra, I = -Vm / Ra
    xo = (np.arange(n) - off) / fs * 1e3
    I[off:] += -(dv / ra * rm / (ra + rm) * 1000.0) * np.exp(-xo[off:] / tau_ms)
    for _ in range(epsc):
        k = rng.integers(off + int(0.1 * fs), n - 400)
        xx = np.arange(400) / fs * 1e3
        I[k:k + 400] += -100 * (1 - np.exp(-xx / 0.5)) * np.exp(-xx / 5.0) / 0.55
    I += rng.normal(0, noise, n)
    V = None
    if command:
        V = np.full(n, hold)
        V[on:off] += dv
        V += rng.normal(0, 0.05, n)
    return t, I, V


def test_command_channel_detection_is_exact_and_measures_dv():
    t, I, V = vc_sweep(dv=-10.0)
    r = qc.measure_sweep(0, t, I, V, QCParams())
    assert r.ok and r.timing_source == "command channel" and r.dv_source == "command channel"
    assert r.onset_ms == pytest.approx(156.3, abs=0.1) and r.offset_ms == pytest.approx(256.3, abs=0.1)
    assert r.dv == pytest.approx(-10.0, abs=0.1)
    assert r.ra == pytest.approx(15.0, rel=0.03) and r.rm == pytest.approx(300.0, rel=0.05)


@pytest.mark.parametrize("fs,onset,pulse,dur", [(20000, 156.3, 100, 2.0), (10000, 53.2, 100, 0.2), (40000, 500.0, 50, 3.0),
                                                 (20000, 20.0, 200, 1.5)])
def test_current_only_detection_any_rate_position_duration(fs, onset, pulse, dur):
    t, I, V = vc_sweep(fs=fs, onset_ms=onset, pulse_ms=pulse, dur_s=dur, command=False)
    r = qc.measure_sweep(0, t, I, V, QCParams())
    assert r.ok and r.timing_source == "current transients" and r.dv_source == "assumed"
    assert r.onset_ms == pytest.approx(onset, abs=0.3) and r.offset_ms == pytest.approx(onset + pulse, abs=0.3)
    assert r.ra == pytest.approx(15.0, rel=0.05)


def test_positive_step_gets_a_positive_ra_with_inferred_sign():
    t, I, V = vc_sweep(dv=+10.0, command=False)
    r = qc.measure_sweep(0, t, I, V, QCParams(v_step=-10.0))
    assert r.ok and r.dv > 0 and r.ra == pytest.approx(15.0, rel=0.05)


def test_step_size_comes_from_the_data_not_the_default():
    t, I, V = vc_sweep(dv=-5.0)
    r = qc.measure_sweep(0, t, I, V, QCParams(v_step=-10.0))
    assert r.dv == pytest.approx(-5.0, abs=0.1) and r.ra == pytest.approx(15.0, rel=0.05)


@pytest.mark.parametrize("ra", [8.0, 20.0, 24.0, 40.0])
def test_ra_recovered_across_values_with_epscs_in_the_trace(ra):
    t, I, V = vc_sweep(ra=ra, rm=250.0, epsc=12, command=False, seed=4)
    r = qc.measure_sweep(0, t, I, V, QCParams())
    assert r.ra == pytest.approx(ra, rel=0.06)


def test_sweep_without_pulse_is_reported_not_guessed():
    t, I, V = vc_sweep()
    flat = np.random.default_rng(1).normal(-9, 3, len(t))
    r = qc.measure_sweep(0, t, flat, None, QCParams())
    assert not r.ok and "no test pulse" in r.note


def test_pulse_late_in_the_sweep_is_found_by_the_whole_sweep_retry():
    t, I, V = vc_sweep(onset_ms=1500.0, dur_s=3.0, command=False)
    r = qc.measure_sweep(0, t, I, V, QCParams(search_ms=500.0))
    assert r.ok and r.onset_ms == pytest.approx(1500.0, abs=0.3)


# ---------------------------------------------------------------------------------------------- verdicts
def fq_with(ras, p=None, **kw):
    p = p or QCParams()
    f = FileQC("x/EPSC/a.abf", "EPSC", "embedded pulse")
    f.sweeps = [SweepRa(i, ok=True, ra=r, rm=200.0, dv=-10.0, dv_source="command channel") for i, r in enumerate(ras)]
    qc._summarise(f, p)
    return f


def test_verdict_is_about_the_limit_and_drift_is_a_separate_flag():
    f = fq_with([10.7, 12.9, 13.5, 13.9])                 # the example cell: well under 25, drifts 30 %
    assert f.status == "PASS" and f.drift_exclude
    assert fq_with([14.0, 14.5, 13.8]).status == "PASS" and not fq_with([14.0, 14.5, 13.8]).drift_exclude
    assert fq_with([24.0, 26.0, 24.5]).status == "WARN"   # one sweep above the limit, mean below
    assert fq_with([26.0, 27.0, 28.0]).status == "FAIL"
    assert fq_with([25.0, 25.0]).status == "FAIL"         # must be BELOW the limit


def test_drift_rule_matches_the_matlab_script_both_directions():
    assert fq_with([10.0, 12.0]).drift_exclude            # >= 1.2 x first
    assert fq_with([10.0, 8.0]).drift_exclude             # <= 0.8 x first
    assert not fq_with([10.0, 11.9, 8.1]).drift_exclude


def test_partial_detection_is_a_warning():
    f = FileQC("a.abf", "EPSC", "embedded pulse")
    f.sweeps = [SweepRa(0, ok=True, ra=12.0, rm=200.0), SweepRa(1, ok=False, note="no test pulse found")]
    qc._summarise(f, QCParams())
    assert f.status == "WARN" and any("not found" in r for r in f.reasons)


# --------------------------------------------------------------------------------------------- discovery
def make_tree(root, files):
    for rel in files:
        p = os.path.join(root, rel)
        os.makedirs(os.path.dirname(p), exist_ok=True)
        open(p, "wb").close()


def test_classify_roles_cells_and_lines(tmp_path):
    root = str(tmp_path)
    make_tree(root, ["line 14/cell 1/EPSC/a.abf", "line 14/cell 1/CM/b.abf", "line 14/cell 1/IO.abf",
                     "line Q2/cell 2/sEPSC.abf", "line Q2/cell 2/membrane.abf", "line Q2/cell 3/epsc rec/c.abf",
                     "line Q2/cell 3/CM test/d.abf", "line 14/cell 1/EPSC/._a.abf", "__MACOSX/line 14/cell 1/EPSC/a.abf",
                     "line 14/cell 4/spontaneous.abf"])
    found, skipped = qc.discover(root)
    got = {os.path.relpath(p, root): (role, cell, line) for p, role, src, cell, line in found}
    assert got["line 14/cell 1/EPSC/a.abf"] == ("EPSC", "line 14/cell 1", "line 14")
    assert got["line 14/cell 1/CM/b.abf"] == ("CM", "line 14/cell 1", "line 14")
    assert got["line Q2/cell 2/sEPSC.abf"] == ("EPSC", "line Q2/cell 2", "line Q2")
    assert got["line Q2/cell 2/membrane.abf"] == ("CM", "line Q2/cell 2", "line Q2")
    assert got["line Q2/cell 3/epsc rec/c.abf"][0] == "EPSC" and got["line Q2/cell 3/CM test/d.abf"][0] == "CM"
    assert got["line Q2/cell 3/epsc rec/c.abf"][1] == got["line Q2/cell 3/CM test/d.abf"][1] == "line Q2/cell 3"
    assert "line 14/cell 1/IO.abf" not in got and "line 14/cell 4/spontaneous.abf" not in got
    assert not any("._a" in k or "__MACOSX" in k for k in got)
    assert skipped == 2


def test_cm_folder_name_is_matched_as_a_word(tmp_path):
    root = str(tmp_path)
    make_tree(root, ["l/c/CM/a.abf", "l/c/cmos_recordings/b.abf", "l/c/Hcmx/c.abf"])
    found, _ = qc.discover(root)
    assert [os.path.basename(p) for p, *_ in found] == ["a.abf"]


# ---------------------------------------------------------------------------------------- real example cell
@pytest.fixture(scope="module")
def mock_tree(cell7_dir, tmp_path_factory):
    root = tmp_path_factory.mktemp("tree")
    for rel, src in (("line 14/cell 1/EPSC/26a01001_sEPSC.abf", "sEPSC.abf"), ("line 14/cell 1/CM/26a01002.abf", "membrane.abf"),
                     ("line Q2/cell 5/EPSC/26a05001.abf", "sEPSC.abf"), ("line Q2/cell 5/IO.abf", "IO.abf")):
        os.makedirs(os.path.dirname(root / rel), exist_ok=True)
        shutil.copy(cell7_dir + src, root / rel)
    return str(root)


def matlab_reference(path, vstep=-10.0):
    """Python port of test_pulsePasca.m (hard-coded 20 kHz indices) for the embedded-pulse file."""
    sweeps, _ = qc.read_sweeps(path)
    out = []
    for t, I, V in sweeps:
        pulse = I[1999:6500]
        base = pulse[:499].mean()
        out.append(vstep / (pulse.min() - base) * 1000)
    return np.array(out)


def test_auto_detection_agrees_with_the_matlab_script_on_the_embedded_pulse(mock_tree):
    path = os.path.join(mock_tree, "line 14/cell 1/EPSC/26a01001_sEPSC.abf")
    ref = matlab_reference(path)
    f = qc.analyse_file(path, "EPSC", "embedded pulse", QCParams())
    ours = np.array([s.ra for s in f.sweeps])
    assert len(ours) == 4 and f.dv_source == "command channel"
    np.testing.assert_allclose(ours, ref, rtol=0.025)              # differences: measured dV (-10.10) and baseline window
    # with the step size forced to the script's -10 mV the agreement is closer
    f10 = qc.analyse_file(path, "EPSC", "embedded pulse", QCParams(min_cmd_mV=1e9, v_step=-10.0))
    np.testing.assert_allclose([s.ra for s in f10.sweeps], ref, rtol=0.01)


def test_cm_file_without_a_recorded_command_channel_uses_the_header_command(mock_tree):
    f = qc.analyse_file(os.path.join(mock_tree, "line 14/cell 1/CM/26a01002.abf"), "CM", "CM file", QCParams())
    assert f.n_ok == f.n_sweeps == 14 and f.dv_source == "command channel" and 9 < f.ra_mean < 12
    assert not any("assumed" in r for r in f.reasons)
    assert f.sweeps[0].onset_ms == pytest.approx(53.3, abs=0.3)


def test_full_scan_report_and_exports(mock_tree, tmp_path):
    rep = qc.scan(mock_tree, QCParams())
    assert len(rep.files) == 3 and len(rep.cells) == 2 and rep.skipped == 1
    assert all(c.status == "PASS" for c in rep.cells) and all(c.drift for c in rep.cells)
    cell1 = [c for c in rep.cells if c.cell == "line 14/cell 1"][0]
    assert len(cell1.files) == 2                                    # EPSC and CM files grouped into one cell
    text = qc.summary_text(rep)
    assert "2 of 2 cells" in text and "drift flag" in text
    x = tmp_path / "ra.xlsx"
    qc.write_xlsx(rep, str(x))
    from openpyxl import load_workbook
    wb = load_workbook(x)
    assert wb.sheetnames == ["Cells", "Files", "Sweeps", "Settings"]
    assert wb["Cells"].max_row == 3 and wb["Files"].max_row == 4 and wb["Sweeps"].max_row == 1 + 4 + 14 + 4
    paths = qc.write_csv(rep, str(tmp_path / "ra"))
    assert all(os.path.exists(p) for p in paths)


def test_limit_can_be_changed(mock_tree):
    rep = qc.scan(mock_tree, QCParams(limit_mohm=12.0))
    assert any(c.status != "PASS" for c in rep.cells)


def test_cli_exit_code_and_output(mock_tree, capsys, tmp_path):
    assert qc.main([mock_tree, "--out", str(tmp_path / "o.xlsx")]) == 0
    assert "Report:" in capsys.readouterr().out
    assert qc.main([mock_tree, "--limit", "5", "--no-xlsx"]) == 1


# ------------------------------------------------------------------------------------------------- GUI tab
def test_qc_tab_scans_and_shows_the_pulse(mock_tree, monkeypatch):
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    pytest.importorskip("PySide6")
    try:
        from PySide6 import QtWidgets
        from ephys_suite.gui import MainWindow
    except Exception as e:                                    # pragma: no cover
        pytest.skip(f"Qt unavailable: {e}")
    app = QtWidgets.QApplication.instance() or QtWidgets.QApplication([])
    w = MainWindow()
    tab = w.qc
    tab.settings.clear()
    tab.lbl_root.setText(mock_tree)
    tab.scan()
    assert tab.tbl.rowCount() == 3 and "2 of 2 cells have Ra" in tab.lbl_sum.text()
    for r in range(3):
        tab.tbl.selectRow(r)
        app.processEvents()
        assert tab.pr.listDataItems()                          # Ra-per-sweep plotted
    tab.chk_bad.setChecked(True)                               # drift-flagged recordings stay listed
    assert tab.tbl.rowCount() == 2
    tab.sp_limit.setValue(11.0)
    tab.scan()
    assert "0 of 2 cells have Ra" in tab.lbl_sum.text() and "2 need attention" in tab.lbl_sum.text()
    w.close()
